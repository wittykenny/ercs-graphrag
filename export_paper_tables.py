# -*- coding: utf-8 -*-
"""Export thesis-ready tables from Project 3 result directories."""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path
from typing import Any, Dict, List

from scripts.ercs_exp_utils import ensure_dir, flatten_metrics, iter_result_dirs, read_config, read_csv_dict, read_json, setup_logger, write_csv, write_json


def summary_rows(root: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for d in iter_result_dirs(root):
        summary = read_json(d / "summary.json", None)
        if summary is None:
            continue
        flat = flatten_metrics("", summary)
        row: Dict[str, Any] = {"result_dir": d.name}
        row.update(flat)
        rows.append(row)
    return rows


def collect_csv(root: Path, filename: str, table_name: str, out_dir: Path) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for d in iter_result_dirs(root):
        p = d / filename
        if not p.exists():
            continue
        for r in read_csv_dict(p):
            item = {"result_dir": d.name}
            item.update(r)
            rows.append(item)
    write_csv(out_dir / f"{table_name}.csv", rows)
    return rows


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default="configs/ercs_hotpotqa_500.yaml")
    ap.add_argument("--root", default=".")
    ap.add_argument("--output-dir", default=None)
    args = ap.parse_args()
    cfg = read_config(args.config)
    out_dir = ensure_dir(args.output_dir or cfg.get("results", {}).get("paper_dir", "results_paper_tables"))
    root = Path(args.root)
    logger = setup_logger("export_paper_tables")

    rows = summary_rows(root)
    write_csv(out_dir / "table_main_results_all_summaries.csv", rows)
    write_json(out_dir / "table_main_results_all_summaries.json", rows)
    logger.info("Exported %d summary rows", len(rows))

    collect_csv(root, "significance_tests.csv", "table_significance_tests", out_dir)
    collect_csv(root, "multi_seed_summary.csv", "table_multi_seed_summary", out_dir)
    collect_csv(root, "community_stats.csv", "table_community_stats", out_dir)
    collect_csv(root, "kg_central_nodes.csv", "table_central_nodes", out_dir)
    collect_csv(root, "triple_quality_sample.csv", "table_triple_quality_samples", out_dir)

    # Copy compact markdown case/failure analyses if available.
    md_dir = ensure_dir(out_dir / "markdown_cases")
    for d in iter_result_dirs(root):
        for name in ["case_studies.md", "failure_cases.md"]:
            src = d / name
            if src.exists():
                shutil.copy2(src, md_dir / f"{d.name}_{name}")

    # Optional Excel workbook if openpyxl is available.
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for csv_path in sorted(out_dir.glob("table_*.csv")):
            ws = wb.create_sheet(csv_path.stem[:31])
            for row in read_csv_dict(csv_path):
                if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
                    ws.append(list(row.keys()))
                ws.append(list(row.values()))
        wb.save(out_dir / "paper_tables.xlsx")
        logger.info("Exported Excel workbook: %s", out_dir / "paper_tables.xlsx")
    except Exception as e:
        logger.warning("Excel export skipped: %s", e)

    logger.info("All paper tables exported to %s", out_dir)


if __name__ == "__main__":
    main()
